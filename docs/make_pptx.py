# Builds docs/Artikel-Trainer-Licensing-and-Mobile-App.pptx.  Usage: pip install python-pptx && python docs/make_pptx.py
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

NAVY = RGBColor(0x14, 0x21, 0x3D)
NAVY2 = RGBColor(0x1F, 0x2E, 0x52)
CREAM = RGBColor(0xFB, 0xFB, 0xF8)
GREY_BG = RGBColor(0xF1, 0xF3, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x24, 0x33)
MUTED = RGBColor(0x4A, 0x55, 0x68)
SOFT = RGBColor(0xD5, 0xDC, 0xEA)
FAINT = RGBColor(0x9A, 0xA6, 0xBF)
BLUE = RGBColor(0x2B, 0x59, 0xC3)
ORANGE = RGBColor(0xE0, 0x7A, 0x1F)
ORANGE_TXT = RGBColor(0xB8, 0x5C, 0x12)
PEACH = RGBColor(0xF2, 0xA6, 0x5A)
BORDER = RGBColor(0xDD, 0xE2, 0xEC)
BAND = RGBColor(0xF7, 0xF8, 0xFB)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
M = 0.89          # side margin (in)
W = 13.333 - 2 * M


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, parts, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """parts: list of paragraphs; each paragraph = (text, size_pt, color, bold) or list of such runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [para]
        for (t, size, color, bold) in runs:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
        p.space_after = Pt(size * 0.35)
    return tb


def box(slide, x, y, w, h, fill, line=None, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def bar(slide, x, y, w, h, color):
    b = box(slide, x, y, w, h, color, radius=False)
    return b


def header(slide, eyebrow, title, dark=False, title_size=32):
    text(slide, M, 0.72, W, 0.4, [(eyebrow.upper(), 14, PEACH if dark else BLUE, True)])
    text(slide, M, 1.08, W, 1.0, [(title, title_size, CREAM if dark else NAVY, True)])


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def footer(slide, t, dark=False):
    text(slide, M, 7.5 - 0.62, W, 0.3, [(t, 12, FAINT if dark else RGBColor(0x5A, 0x64, 0x78), False)])


def card(slide, x, y, w, h, title, body, num=None, top_color=None, fill=WHITE, line=BORDER,
         title_color=NAVY, body_color=MUTED, label=None, label_color=BLUE):
    box(slide, x, y, w, h, fill, line)
    if top_color:
        bar(slide, x, y, w, 0.08, top_color)
    cy = y + 0.3
    parts = []
    if num:
        parts.append((num, 32, BLUE, True))
    if label:
        parts.append((label.upper(), 13, label_color, True))
    parts.append((title, 18, title_color, True))
    parts.append((body, 14, body_color, False))
    text(slide, x + 0.3, cy, w - 0.6, h - 0.5, parts)


def table(slide, y, widths, rows, font=13):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(M), Inches(y), Inches(W), Inches(0.4 * len(rows)))
    tbl = shape.table
    for j, wpct in enumerate(widths):
        tbl.columns[j].width = Emu(int(Inches(W) * wpct))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (BAND if i % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.08)
            bold = i == 0 or val.startswith("**")
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val.strip("*")
            r.font.name = FONT
            r.font.size = Pt(font)
            r.font.bold = bold
            r.font.color.rgb = CREAM if i == 0 else INK
    return tbl


# 1 Cover -------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, M, 2.1, W, 3.6, [
    ("DER · DIE · DAS", 16, PEACH, True),
    ("Anshi German Learning App", 48, CREAM, True),
    ("Licensing check and the road to a mobile app", 24, SOFT, False),
])
footer(s, "A practical summary, not legal advice · September 2026", dark=True)
notes(s, "Two questions: can we publish the tool, and how do we turn it into an app. Short answers first, then the details.")

# 2 Short answer ------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, CREAM)
header(s, "Can I publish it?", "Yes, for free or even commercially.", title_size=40)
text(s, M, 1.95, W, 0.5, [('No part of the app is "all rights reserved". Three conditions apply:', 18, MUTED, False)])
cw = (W - 2 * 0.22) / 3
for i, (t, b) in enumerate([
        ("Credit the open sources", "Wiktionary, Tatoeba, the frequency lists, fonts and libraries."),
        ("Keep Wiktionary data open", "Data files built from Wiktionary stay under CC BY-SA if you share them."),
        ("Add two pages", "An About / Credits page and a privacy policy.")]):
    card(s, M + i * (cw + 0.22), 2.9, cw, 3.2, t, b, num=str(i + 1))
notes(s, "This is a practical summary, not legal advice. For a commercial launch, a quick check with someone who knows open licences is worth it.")

# 3 Data --------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, GREY_BG)
header(s, "What's inside · 1 of 2", "Data from open sources")
table(s, 2.05, [0.31, 0.22, 0.17, 0.30], [
    ["Part of the app", "Source", "Licence", "What you must do"],
    ["Extra nouns and verbs, verb forms, adjective meanings, plurals, syllable splits", "Wiktionary (via kaikki.org, gambolputty)", "**CC BY-SA", "Credit Wiktionary; share these data files under CC BY-SA too"],
    ["17,236 example sentences", "Tatoeba", "**CC BY 2.0 FR", "Credit Tatoeba (the per-sentence links already do)"],
    ["Frequency ranking (Top 5000)", "Leipzig Corpora, OpenSubtitles word list", "Statistics; sources ask for credit", "Name both as sources"],
], font=14)
footer(s, "Rankings are computed statistics, but crediting the corpora is the safe and polite choice.")
notes(s, "Wiktionary is the only source with a ShareAlike condition, explained on slide 5.")

# 4 Code --------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, GREY_BG)
header(s, "What's inside · 2 of 2", "Your own work, tools and fonts")
table(s, 2.05, [0.33, 0.20, 0.19, 0.28], [
    ["Part of the app", "Source", "Licence", "What you must do"],
    ["App code, design, themed noun lists, small-word English, picture choices", "This project", "**Yours", "Nothing; choose any licence"],
    ["Excel library", "SheetJS (xlsx)", "**Apache 2.0", "Include its licence text"],
    ["Hyphenation patterns", "TeX hyphenation project", "**MIT", "Keep the notice (already in the file)"],
    ["Fonts: Playwrite DE, Nunito", "Google Fonts", "**SIL OFL", "Free to bundle; include the OFL notice"],
    ["Emoji pictures and voices", "The user's device", "Not shipped", "Nothing"],
], font=14)
notes(s, "Emoji are drawn by the phone's own emoji font and voices come from the phone's text-to-speech, so neither is part of the app.")

# 5 ShareAlike --------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, CREAM)
header(s, "ShareAlike, explained", "Each part keeps its own licence")
cw = (W - 0.22) / 2
for i, (t, sub, items, col) in enumerate([
        ("Must stay CC BY-SA", "Data files built from Wiktionary:", ["09-extended-nouns.js", "verbs.js, words.js", "plurals.js, syllables.js"], ORANGE),
        ("Your choice of licence", "Everything written for this project:", ["app.js, index.html, style.css", "Themed noun lists (01–08)", "pictures.js"], BLUE)]):
    x = M + i * (cw + 0.22)
    box(s, x, 2.05, cw, 3.9, WHITE, BORDER)
    bar(s, x, 2.05, cw, 0.08, col)
    text(s, x + 0.35, 2.4, cw - 0.7, 3.4, [(t, 20, NAVY, True), (sub, 14, MUTED, False)] + [("•  " + it, 15, INK, False) for it in items])
footer(s, "An app that bundles parts under different licences is fine, as long as each part's terms are followed.")
notes(s, "ShareAlike means that anyone who receives the Wiktionary-derived data files may reuse them under the same licence. It does not force the app code to be open source.")

# 6 Checklist ---------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, GREY_BG)
header(s, "Before publishing", "Six small tasks")
items = [("About / Credits page", "Lists every source and licence."),
         ("Privacy policy", "Short: no data is collected."),
         ("LICENSE and NOTICE files", "Your licence plus third-party notices."),
         ("Check the app name", 'Search the stores for "Anshi German Learning App" first.'),
         ("Skim the examples", 'Top 5000, or add a "report this sentence" button.'),
         ("Bundle the fonts", "So cursive works offline.")]
cw = (W - 2 * 0.2) / 3
for i, (t, b) in enumerate(items):
    r, c = divmod(i, 3)
    card(s, M + c * (cw + 0.2), 2.05 + r * 2.15, cw, 1.95, t, b, label="✓ Task " + str(i + 1))
notes(s, "The example sentences are user-written on Tatoeba. They were filtered for kid safety, but nobody has read all 17,000.")

# 7 Kids --------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, CREAM)
text(s, M, 1.9, 5.3, 4, [("KIDS AND PRIVACY", 14, BLUE, True), ("No data leaves the device.", 40, NAVY, True),
                         ("Settings, scores and added words are saved in the browser on the device only.", 16, MUTED, False)])
for i, (t, b) in enumerate([
        ("No ads, analytics or accounts", "Keep it that way to qualify for kids' programmes."),
        ("Store rules for kids", 'Apple Kids category and Google "Designed for Families": no tracking, and a parental gate before external links such as Tatoeba.'),
        ("Privacy policy still required", "Both stores want a link, even when nothing is collected.")]):
    y = 1.2 + i * 1.75
    box(s, 6.9, y, 5.55, 1.55, WHITE)
    bar(s, 6.9, y, 0.08, 1.55, BLUE)
    text(s, 7.2, y + 0.22, 5.0, 1.2, [(t, 17, NAVY, True), (b, 13, MUTED, False)])
notes(s, "The parental gate matters for the Tatoeba source links in the dictionary: in the Kids category, links that leave the app need one.")

# 8 Part 2 ------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, M, 2.2, 10.5, 3.5, [("PART 2", 16, PEACH, True), ("Making it an app", 56, CREAM, True),
                            ("About 5 MB of offline data, with pronunciation at the heart of it: the voice decides which route works best.", 22, SOFT, False)])
notes(s, "Pronunciation now drives much of the app: words, syllables and example sentences. That is why text-to-speech support shapes the recommendation.")

# 9 Options -----------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, CREAM)
header(s, "The options", "Four ways to get onto phones")
table(s, 2.05, [0.25, 0.22, 0.23, 0.30], [
    ["Option", "Effort and cost", "Voice (text-to-speech)", "Notes"],
    ["**1. Installable web app (PWA)", "About 1 hour; free hosting", "Works (real browser)", "Home-screen icon, offline; not in the stores"],
    ["**2. Play Store via PWABuilder", "A few hours; $25 once", "Works (runs in Chrome)", "Built from the PWA; updates go live without a new release"],
    ["**3. Capacitor app (iOS + Android)", "2–4 days; Apple $99/year + a Mac", "iOS works; Android needs a plugin", "Plugins for downloads and sharing; printing differs"],
    ["**Rewrite in Flutter or React Native", "Weeks", "Native", "Not worth it: everything would be rebuilt"],
], font=14)
notes(s, "Option 2 uses a Trusted Web Activity: the PWA shown full screen inside Chrome, packaged for the Play Store.")

# 10 Voice ------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, GREY_BG)
header(s, "Why the voice matters", "Where the built-in voice works")
cw = (W - 2 * 0.22) / 3
for i, (lab, lc, top, t, b) in enumerate([
        ("Works", BLUE, BLUE, "Browser, PWA and Play Store (PWABuilder)", "Words, syllables and sentences play unchanged."),
        ("Works", BLUE, BLUE, "iPhone app (Capacitor on iOS)", "The iOS web view includes the speech engine."),
        ("Needs a plugin", ORANGE_TXT, ORANGE, "Android app (Capacitor on Android)", "Android's built-in web view has no text-to-speech; add a native speech plugin.")]):
    card(s, M + i * (cw + 0.22), 2.05, cw, 3.0, t, b, top_color=top, label=lab, label_color=lc, line=None)
text(s, M, 5.45, W, 0.6, [("So for Android, the Play Store route through PWABuilder keeps the voice working with no extra code.", 16, MUTED, False)])
notes(s, "The same applies to file downloads: the Excel template and export need the share sheet inside a Capacitor app.")

# 11 Costs ------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, CREAM)
header(s, "Costs and requirements", "What each step costs")
cw = (W - 2 * 0.22) / 3
for i, (amt, col, t, b) in enumerate([
        ("$0", BLUE, "Web app (PWA)", "Free hosting on GitHub Pages, Netlify or Cloudflare Pages."),
        ("$25", BLUE, "Google Play, one-time", "Developer account; any computer works."),
        ("$99", ORANGE_TXT, "Apple, per year", "Developer account plus a Mac, or a cloud build service.")]):
    x = M + i * (cw + 0.22)
    box(s, x, 2.2, cw, 3.4, WHITE, BORDER)
    text(s, x + 0.35, 2.5, cw - 0.7, 3.0, [(amt, 54, col, True), (t, 18, NAVY, True), (b, 14, MUTED, False)])
notes(s, 'Both stores also need a privacy policy link. Apple sometimes rejects apps that are "just a website"; offline use, pronunciation and file import count as app features.')

# 12 Roadmap ----------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, GREY_BG)
header(s, "Recommendation", "Three steps, cheapest first")
aw, gap = 0.45, 0.15
cw = (W - 2 * (aw + 2 * gap)) / 3
for i, (t, b) in enumerate([
        ("Installable web app", "Offline caching, local fonts, icons, About / Credits tab, privacy page. Works on every phone right away."),
        ("Android on Google Play", "Package the web app with PWABuilder. Voice works unchanged. $25 once."),
        ("iPhone App Store", 'Wrap with Capacitor when you want to be listed. Meanwhile: "Add to Home Screen".')]):
    x = M + i * (cw + aw + 2 * gap)
    card(s, x, 2.1, cw, 3.4, t, b, num=str(i + 1), line=None)
    if i < 2:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw + gap), Inches(3.65), Inches(aw), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = FAINT; a.line.fill.background()
notes(s, "Step 1 is the foundation for both store versions, so no work is wasted.")

# 13 Next -------------------------------------------------------------
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, M, 1.7, 5.4, 4.5, [("NEXT STEPS", 14, PEACH, True), ("Three decisions, then step 1", 40, CREAM, True),
                           ("Step 1 covers the web app setup, local fonts, the About / Credits tab, the privacy policy and the LICENSE / NOTICE files.", 16, SOFT, False)])
for i, (t, b) in enumerate([
        ("Publisher name", "Shown in the stores and on the Credits page."),
        ("Licence for your own code", 'Open source (e.g. MIT) or "all rights reserved".'),
        ("App name", 'Decided: "Anshi German Learning App".')]):
    y = 1.3 + i * 1.7
    box(s, 6.9, y, 5.55, 1.5, NAVY2)
    text(s, 7.2, y + 0.25, 5.0, 1.1, [(t, 18, CREAM, True), (b, 14, SOFT, False)])
notes(s, "Once these three are decided, step 1 can start right away.")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Artikel-Trainer-Licensing-and-Mobile-App.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("saved", out, len(prs.slides), "slides")
